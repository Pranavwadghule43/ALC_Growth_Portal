"""Phase 3B canonical RCU Pune master builder: DCU/SBU normalisation, ALC Code text handling,
row checks, confirmed counts, explicit Active default, and acceptance of the output by the
Phase 3A validator. The builder never uses a database.

Self-contained: the source spreadsheets are generated in a temp directory by
``tests.rcu_fixtures`` (real source spellings, numeric code cells, confirmed per-SBU counts).
The git-ignored real files are only used by the optional checks at the end, which are skipped
when those files are not present (e.g. on a fresh clone)."""

import argparse
import csv
import io
from pathlib import Path

import openpyxl
import pytest

from app.services import master_import
from scripts import build_rcu_master as builder
from scripts.build_rcu_master import SourceRow
from tests.rcu_fixtures import SOURCE_LAYOUT, source_rows, write_source_workbooks

DATA = Path(__file__).resolve().parents[3] / "data"
NASHIK_MASTER = DATA / "ALC-MASTER.csv"  # committed Nashik master (tracked in git)
# Local, git-ignored real master data: only used by the optional checks at the end.
REAL_SOURCES = [DATA / "Ahilya_Nagar.xlsx", DATA / "Pune.xlsx"]
REAL_NEW_MASTER = DATA / "RCU-PUNE-NEW-ALCS.csv"
requires_real_data = pytest.mark.skipif(
    not all(p.exists() for p in [*REAL_SOURCES, REAL_NEW_MASTER]),
    reason="local real master data (git-ignored) is not present",
)


def row(
    code="17210005",
    name="Balaji Computer Education",
    dcu="Ahilyanagar",
    sbu="Ahilyanagar_sbu1",
    n=2,
):
    return SourceRow(source="t.xlsx", row=n, alc_code=code, alc_name=name, dcu=dcu, sbu=sbu)


def small(*rows, active="yes"):
    """Build a handful of rows; per-SBU confirmed counts are not the point here."""
    return builder.build(list(rows), active, allow_count_mismatch=True)


def row_errors(result) -> str:
    return "\n".join(e for e in result.errors if "count mismatch" not in e)


def xlsx(*rows, number_format=None) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["ALC Code", "ALC Name", "DCU", "SBU"])
    for values in rows:
        sheet.append(list(values))
    if number_format:
        for cells in sheet.iter_rows(min_row=2, max_col=1):
            cells[0].number_format = number_format
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def sources(tmp_path_factory):
    """Synthetic ``Ahilya_Nagar.xlsx`` and ``Pune.xlsx`` in a module temp directory."""
    return write_source_workbooks(tmp_path_factory.mktemp("rcu-sources"))


@pytest.fixture(scope="module")
def built(sources):
    return builder.build([r for path in sources for r in builder.read_source(path)], "yes")


# --------------------------------------------------------------------------- #
# Normalisation
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("source", ["Ahilyanagar", "ahilyanagar ", "Ahilya Nagar"])
def test_ahilya_nagar_dcu_normalisation(source):
    result = small(row(dcu=source))
    assert result.rows[0]["DCU"] == "Ahilya Nagar"
    assert not row_errors(result)


def test_ahilyanaga_sbu10_spelling_normalised():
    result = small(row(sbu="Ahilyanaga_sbu10"))
    assert result.rows[0]["SBU"] == "Ahilyanagar_sbu10"
    assert any("'Ahilyanaga_sbu10' -> 'Ahilyanagar_sbu10'" in n for n in result.notes)


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5])
def test_pune_north_sbu_mapping_kept(n):
    result = small(row(dcu="Pune North", sbu=f"SBU_Pune_North_{n}"))
    assert (result.rows[0]["DCU"], result.rows[0]["SBU"]) == ("Pune North", f"SBU_Pune_North_{n}")


@pytest.mark.parametrize("variant", ["Pune North 1", "SBU Pune North 1", "SBU_Pune_North_6"])
def test_pune_north_unconfirmed_spellings_are_unresolved(variant):
    result = small(row(dcu="Pune North", sbu=variant))
    assert f"unresolved SBU '{variant}'" in row_errors(result)
    assert result.csv_bytes == b""


@pytest.mark.parametrize(
    "coordinator,sbu",
    [
        ("Bhagyashree Gaikwad", "pune_south_sbu_4"),
        ("Ajinkya Chavan", "pune_south_sbu_2"),
        ("Aniket Marne", "pune_south_sbu_3"),
    ],
)
def test_pune_south_coordinator_names_map_to_sbu_ids(coordinator, sbu):
    result = small(row(dcu="Pune South", sbu=coordinator))
    assert result.rows[0]["SBU"] == sbu
    assert coordinator not in result.csv_bytes.decode()  # never stored as an SBU id
    assert any(f"'{coordinator}' -> '{sbu}'" in n for n in result.notes)


def test_unknown_dcu_and_coordinator_are_unresolved():
    result = small(
        row(dcu="Satara", n=2),
        row(code="87219999", dcu="Pune South", sbu="Someone Else", n=3),
    )
    errors = row_errors(result)
    assert "unresolved DCU 'Satara'" in errors
    assert "unresolved SBU 'Someone Else' for DCU 'Pune South'" in errors


def test_nashik_rows_rejected_unless_included():
    result = small(row(dcu="Nashik", sbu="SBU 4"))
    assert "Nashik is already imported" in row_errors(result)


# --------------------------------------------------------------------------- #
# ALC Code / Name handling
# --------------------------------------------------------------------------- #
def test_numeric_code_cell_becomes_integer_text():
    (source,) = builder.read_xlsx(
        xlsx([17210005.0, "Name", "Ahilyanagar", "Ahilyanagar_sbu1"]), "t.xlsx"
    )
    assert source.alc_code == "17210005" and isinstance(source.alc_code, str)
    assert source.numeric_code
    result = small(source)
    assert b"17210005.0" not in result.csv_bytes and b",17210005," in result.csv_bytes


def test_leading_zeros_preserved():
    text_cell, padded = builder.read_xlsx(
        xlsx(["00123456", "Text Code", "Pune North", "SBU_Pune_North_1"]), "t.xlsx"
    ) + builder.read_xlsx(
        xlsx([123456, "Padded", "Pune North", "SBU_Pune_North_1"], number_format="00000000"),
        "t.xlsx",
    )
    assert text_cell.alc_code == "00123456" and not text_cell.numeric_code
    assert padded.alc_code == "00123456"
    (from_csv,) = builder.read_csv(
        b"ALC Code,ALC Name,DCU,SBU\n0012345,Csv Code,Pune North,SBU_Pune_North_2\n", "t.csv"
    )
    assert from_csv.alc_code == "0012345"
    result = small(text_cell, from_csv)
    written = list(csv.DictReader(io.StringIO(result.csv_bytes.decode())))
    assert [r["ALC Code"] for r in written] == ["00123456", "0012345"]
    (parsed, _) = master_import.parse_source("m.csv", result.csv_bytes)
    assert parsed.alc_code == "00123456"


def test_fractional_numeric_code_rejected():
    (source,) = builder.read_xlsx(
        xlsx([1721.5, "Frac", "Pune North", "SBU_Pune_North_1"]), "t.xlsx"
    )
    assert "not a whole number" in row_errors(small(source))


def test_alc_name_copied_exactly():
    name = "Balaji Computer Education, Kedgaon"
    assert small(row(name=name)).rows[0]["ALC Name"] == name


def test_duplicate_alc_code_rejected():
    result = small(row(name="One", n=2), row(name="Two", sbu="Ahilyanagar_sbu2", n=3))
    errors = row_errors(result)
    assert "duplicate ALC Code '17210005' (also t.xlsx row 3)" in errors
    assert "duplicate ALC Code '17210005' (also t.xlsx row 2)" in errors
    assert result.csv_bytes == b""


def test_duplicate_full_row_rejected():
    result = small(row(n=2), row(n=3))
    assert "duplicate row (identical to t.xlsx row 2)" in row_errors(result)


def test_blank_alc_code_rejected():
    assert "blank ALC Code" in row_errors(small(row(code="")))


def test_blank_alc_name_rejected():
    assert "blank ALC Name" in row_errors(small(row(name="  ")))


# --------------------------------------------------------------------------- #
# Active default
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("flag,value", [("yes", "Yes"), ("no", "No"), ("YES", "Yes")])
def test_explicit_active_default(flag, value):
    result = small(row(), active=flag)
    assert result.rows[0]["Active"] == value and result.active == value
    assert any("--default-active" in n and "NOT from the source" in n for n in result.notes)


@pytest.mark.parametrize("flag", ["", None, "maybe"])
def test_active_default_is_required(flag):
    with pytest.raises(ValueError):
        builder.build([row()], flag)


# --------------------------------------------------------------------------- #
# Full-size representative sources (synthetic, confirmed counts)
# --------------------------------------------------------------------------- #
def test_sources_build_cleanly(built):
    assert built.ok, built.errors[:5]
    assert not built.warnings


def test_expected_ahilya_nagar_counts(built):
    assert built.counts["Ahilya Nagar"] == {
        "Ahilyanagar_sbu1": 58,
        "Ahilyanagar_sbu2": 49,
        "Ahilyanagar_sbu5": 51,
        "Ahilyanagar_sbu10": 45,
    }
    assert sum(built.counts["Ahilya Nagar"].values()) == 203


def test_expected_pune_north_counts(built):
    assert built.counts["Pune North"] == {
        "SBU_Pune_North_1": 53,
        "SBU_Pune_North_2": 48,
        "SBU_Pune_North_3": 55,
        "SBU_Pune_North_4": 52,
        "SBU_Pune_North_5": 55,
    }
    assert sum(built.counts["Pune North"].values()) == 263


def test_expected_pune_south_counts(built):
    assert built.counts["Pune South"] == {
        "pune_south_sbu_4": 46,
        "pune_south_sbu_2": 42,
        "pune_south_sbu_3": 31,
    }
    assert sum(built.counts["Pune South"].values()) == 119


def test_total_new_count_is_585(built):
    assert len(built.rows) == 585
    assert sum(sum(s.values()) for s in built.counts.values()) == 585
    assert "Nashik" not in built.counts


def test_every_row_mapped_to_its_canonical_dcu_and_sbu(built):
    expected = [(r["dcu"], r["sbu"]) for f in SOURCE_LAYOUT for r in source_rows(f)]
    assert [(r["DCU"], r["SBU"]) for r in built.rows] == expected


def test_normalisation_notes_report_counts(built):
    notes = "\n".join(built.notes)
    assert "DCU 'Ahilyanagar' -> 'Ahilya Nagar': 203 rows" in notes
    assert "SBU 'Ahilyanaga_sbu10' -> 'Ahilyanagar_sbu10' (Ahilya Nagar): 45 rows" in notes
    assert "SBU 'Bhagyashree Gaikwad' -> 'pune_south_sbu_4' (Pune South): 46 rows" in notes
    assert "SBU 'Ajinkya Chavan' -> 'pune_south_sbu_2' (Pune South): 42 rows" in notes
    assert "SBU 'Aniket Marne' -> 'pune_south_sbu_3' (Pune South): 31 rows" in notes
    assert "585 ALC Codes were numeric spreadsheet cells" in notes


def test_canonical_six_column_output(built):
    reader = csv.reader(io.StringIO(built.csv_bytes.decode()))
    header = next(reader)
    assert header == ["RCU", "DCU", "SBU", "ALC Code", "ALC Name", "Active"]
    body = list(reader)
    assert len(body) == 585 and all(len(r) == 6 for r in body)
    assert {r[0] for r in body} == {"RCU Pune"}
    assert {r[1] for r in body} == {"Ahilya Nagar", "Pune North", "Pune South"}
    assert {r[5] for r in body} == {"Yes"}
    assert not any(r[3].endswith(".0") for r in body)
    for coordinator in ("Bhagyashree Gaikwad", "Ajinkya Chavan", "Aniket Marne"):
        assert coordinator not in built.csv_bytes.decode()


def test_source_codes_and_names_preserved_exactly(built, sources):
    source = []
    for path in sources:
        sheet = openpyxl.load_workbook(path, data_only=True).active
        source += [(str(int(c)), n) for c, n, *_ in sheet.iter_rows(min_row=2, values_only=True)]
    assert [(r["ALC Code"], r["ALC Name"]) for r in built.rows] == source
    assert any("," in name for _, name in source)  # quoted names survive the round trip


def test_output_accepted_by_phase3a_file_validator(built):
    records = master_import.parse_source("m.csv", built.csv_bytes)
    master_import.check_file(records)
    assert len(records) == 585
    assert [r.errors for r in records if r.errors] == []


def test_builder_output_is_deterministic(built, sources):
    again = builder.build([r for path in sources for r in builder.read_source(path)], "yes")
    assert again.csv_bytes == built.csv_bytes


def test_complete_master_with_nashik_is_784(sources):
    rows = [r for path in sources for r in builder.read_source(path)]
    rows += builder.read_nashik(NASHIK_MASTER)
    result = builder.build(rows, "yes", include_nashik=True)
    assert result.ok, result.errors[:5]
    assert result.counts["Nashik"] == {"SBU 4": 71, "SBU 6": 58, "SBU 7": 70}
    assert len(result.rows) == 784


def test_count_mismatch_blocks_unless_allowed(sources):
    rows = [r for r in builder.read_source(sources[0])][:-1]  # drop one Ahilya Nagar ALC
    blocked = builder.build(rows, "yes")
    assert not blocked.ok and blocked.csv_bytes == b""
    assert any("count mismatch" in e for e in blocked.errors)
    allowed = builder.build(rows, "yes", allow_count_mismatch=True)
    assert allowed.ok and any("count mismatch" in w for w in allowed.warnings)


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def cli_args(tmp_path, sources, **extra):
    return argparse.Namespace(
        sources=sources,
        default_active=extra.get("default_active", "yes"),
        out=tmp_path / "out.csv",
        include_nashik=extra.get("include_nashik"),
        allow_count_mismatch=extra.get("allow_count_mismatch", False),
    )


def test_cli_refuses_to_write_when_errors_remain(tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text("ALC Code,ALC Name,DCU,SBU\n,No Code,Pune North,SBU_Pune_North_1\n")
    args = cli_args(tmp_path, [bad], allow_count_mismatch=True)
    assert builder.run(args) == 1
    assert not args.out.exists()


def test_cli_writes_valid_master(tmp_path, sources, built):
    args = cli_args(tmp_path, sources)
    assert builder.run(args) == 0
    assert args.out.read_bytes() == built.csv_bytes


def test_cli_requires_default_active(monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", ["build_rcu_master", "Ahilya_Nagar.xlsx"])
    with pytest.raises(SystemExit) as exc:
        builder.main()
    assert exc.value.code == 2
    assert "--default-active" in capsys.readouterr().err


# --------------------------------------------------------------------------- #
# Optional: the real, git-ignored source files (skipped when absent)
# --------------------------------------------------------------------------- #
@requires_real_data
def test_real_sources_match_confirmed_counts_and_local_master():
    result = builder.build([r for p in REAL_SOURCES for r in builder.read_source(p)], "yes")
    assert result.ok, result.errors[:5]
    assert {d: sum(c.values()) for d, c in result.counts.items()} == {
        "Ahilya Nagar": 203,
        "Pune North": 263,
        "Pune South": 119,
    }
    assert REAL_NEW_MASTER.read_bytes() == result.csv_bytes
